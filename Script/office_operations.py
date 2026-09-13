from __future__ import annotations

import os
import platform
import shutil
import subprocess
import tempfile
from typing import Callable, Optional

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from Config import config

ProgressCallback = Optional[Callable[[int, int, str], None]]
# callback(current_step, total_steps, message) -> None


def _libreoffice_convert_to_pdf(input_path: str, output_path: str) -> str:
    """Shared LibreOffice-headless conversion, used for both Word->PDF and
    PowerPoint->PDF — LibreOffice handles either input format the same way."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice isn't installed on this computer.")

    with tempfile.TemporaryDirectory() as tmp_dir:
        result = subprocess.run(
            [soffice, "--headless", "--norestore", "--convert-to", "pdf",
             "--outdir", tmp_dir, input_path],
            capture_output=True, text=True, timeout=120,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or "conversion failed")

        base_name = os.path.splitext(os.path.basename(input_path))[0]
        produced = os.path.join(tmp_dir, f"{base_name}.pdf")
        if not os.path.exists(produced):
            raise RuntimeError("LibreOffice did not produce an output file")

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        shutil.move(produced, output_path)

    return output_path

EMU_PER_INCH = 914400  # python-pptx works in English Metric Units


class PdfToWordConverter:
    """Converts a PDF into an editable .docx using pdf2docx (pure Python,
    no external engine needed) — reconstructs real text/paragraphs and
    embeds images, rather than just rendering pages as pictures. Complex
    layouts (multi-column, unusual tables) may not reconstruct perfectly,
    but plain text-and-image documents convert cleanly."""

    @staticmethod
    def convert(input_path: str, output_path: str,
                progress: ProgressCallback = None) -> str:
        import logging

        from pdf2docx import Converter

        # pdf2docx logs a lot of INFO-level progress lines via the root
        # logger by default — quiet that down for a GUI app with no console.
        logging.getLogger().setLevel(logging.WARNING)

        if progress:
            progress(0, 1, "Converting...")

        cv = Converter(input_path)
        try:
            cv.convert(output_path)
        finally:
            cv.close()

        if progress:
            progress(1, 1, "Done")
        return output_path


class WordToPdfConverter:
    """Converts a .docx file to .pdf, using MS Word if available, or
    LibreOffice headless as a fallback."""

    @staticmethod
    def convert(input_path: str, output_path: str,
                progress: ProgressCallback = None) -> str:
        if progress:
            progress(0, 1, "Converting...")

        errors = []
        if platform.system() in ("Windows", "Darwin"):
            try:
                return WordToPdfConverter._convert_with_word(
                    input_path, output_path, progress
                )
            except Exception as e:  # noqa: BLE001
                errors.append(f"MS Word: {e}")

        if shutil.which("soffice") or shutil.which("libreoffice"):
            try:
                return WordToPdfConverter._convert_with_libreoffice(
                    input_path, output_path, progress
                )
            except Exception as e:  # noqa: BLE001
                errors.append(f"LibreOffice: {e}")

        detail = " / ".join(errors) if errors else "no converter available"
        raise RuntimeError(
            "Couldn't convert this file to PDF. Word-to-PDF needs either "
            f"Microsoft Word or LibreOffice installed on this computer ({detail})."
        )

    @staticmethod
    def _convert_with_word(input_path: str, output_path: str,
                            progress: ProgressCallback) -> str:
        from docx2pdf import convert as docx2pdf_convert
        docx2pdf_convert(input_path, output_path)
        if progress:
            progress(1, 1, "Done")
        return output_path

    @staticmethod
    def _convert_with_libreoffice(input_path: str, output_path: str,
                                   progress: ProgressCallback) -> str:
        result = _libreoffice_convert_to_pdf(input_path, output_path)
        if progress:
            progress(1, 1, "Done")
        return result


class PptToPdfConverter:
    """Converts a .pptx file to .pdf via LibreOffice headless (no MS
    PowerPoint dependency needed)."""

    @staticmethod
    def convert(input_path: str, output_path: str,
                progress: ProgressCallback = None) -> str:
        if progress:
            progress(0, 1, "Converting...")
        if not (shutil.which("soffice") or shutil.which("libreoffice")):
            raise RuntimeError(
                "PowerPoint-to-PDF needs LibreOffice installed on this computer."
            )
        result = _libreoffice_convert_to_pdf(input_path, output_path)
        if progress:
            progress(1, 1, "Done")
        return result


class PdfToPptConverter:
    """Converts a PDF to .pptx, either as full-slide images (fast, exact
    visual copy) or as real editable text boxes + images (approximate
    layout, but every word is selectable/editable in PowerPoint)."""

    @staticmethod
    def convert(input_path: str, output_path: str, editable: bool = False,
                progress: ProgressCallback = None) -> str:
        if editable:
            return PdfToPptConverter._convert_editable(input_path, output_path, progress)
        return PdfToPptConverter._convert_image_based(input_path, output_path, progress)

    @staticmethod
    def _convert_image_based(input_path: str, output_path: str,
                              progress: ProgressCallback = None) -> str:
        with pymupdf.open(input_path) as doc:
            total = doc.page_count
            if total == 0:
                raise ValueError("This PDF has no pages.")

            first_rect = doc[0].rect
            slide_w_in = first_rect.width / 72.0
            slide_h_in = first_rect.height / 72.0

            prs = Presentation()
            prs.slide_width = Emu(int(slide_w_in * EMU_PER_INCH))
            prs.slide_height = Emu(int(slide_h_in * EMU_PER_INCH))
            blank_layout = prs.slide_layouts[6]  # "Blank"

            with tempfile.TemporaryDirectory() as tmp_dir:
                for i in range(total):
                    if progress:
                        progress(i + 1, total, f"Rendering slide {i + 1}/{total}")
                    page = doc[i]
                    mat = pymupdf.Matrix(
                        config.PPTX_RENDER_SCALE, config.PPTX_RENDER_SCALE
                    )
                    pix = page.get_pixmap(matrix=mat)
                    img_path = os.path.join(tmp_dir, f"page_{i + 1}.png")
                    pix.save(img_path)

                    slide = prs.slides.add_slide(blank_layout)
                    slide.shapes.add_picture(
                        img_path, 0, 0, width=prs.slide_width, height=prs.slide_height
                    )

                prs.save(output_path)

        return output_path

    @staticmethod
    def _convert_editable(input_path: str, output_path: str,
                           progress: ProgressCallback = None) -> str:
        """Rebuilds each page from its actual text spans (as real, editable
        PowerPoint text boxes matching position/size/bold/italic/color as
        closely as PowerPoint allows) and its embedded images (placed at
        their original position). Vector drawings/shapes aren't
        reconstructed — this mode trades exact visual fidelity for having
        genuinely editable text."""
        with pymupdf.open(input_path) as doc:
            total = doc.page_count
            if total == 0:
                raise ValueError("This PDF has no pages.")

            first_rect = doc[0].rect
            slide_w_in = first_rect.width / 72.0
            slide_h_in = first_rect.height / 72.0

            prs = Presentation()
            prs.slide_width = Emu(int(slide_w_in * EMU_PER_INCH))
            prs.slide_height = Emu(int(slide_h_in * EMU_PER_INCH))
            blank_layout = prs.slide_layouts[6]

            with tempfile.TemporaryDirectory() as tmp_dir:
                for i in range(total):
                    if progress:
                        progress(i + 1, total, f"Rebuilding slide {i + 1}/{total}")
                    page = doc[i]
                    slide = prs.slides.add_slide(blank_layout)

                    PdfToPptConverter._place_images(doc, page, slide, tmp_dir, i)
                    PdfToPptConverter._place_text(page, slide)

                prs.save(output_path)

        return output_path

    @staticmethod
    def _place_images(doc, page, slide, tmp_dir, page_index):
        for img_i, img in enumerate(page.get_images(full=True)):
            xref = img[0]
            try:
                rects = page.get_image_rects(xref)
            except Exception:
                rects = []
            if not rects:
                continue
            try:
                base = doc.extract_image(xref)
            except Exception:
                continue
            ext = base.get("ext", "png")
            img_path = os.path.join(tmp_dir, f"p{page_index}_img{img_i}.{ext}")
            with open(img_path, "wb") as f:
                f.write(base["image"])
            for r in rects:
                left = Emu(int(r.x0 / 72 * EMU_PER_INCH))
                top = Emu(int(r.y0 / 72 * EMU_PER_INCH))
                width = Emu(max(int((r.x1 - r.x0) / 72 * EMU_PER_INCH), 1))
                height = Emu(max(int((r.y1 - r.y0) / 72 * EMU_PER_INCH), 1))
                try:
                    slide.shapes.add_picture(img_path, left, top, width=width, height=height)
                except Exception:
                    pass  # skip images python-pptx can't embed (unsupported format, etc.)

    @staticmethod
    def _place_text(page, slide):
        text_dict = page.get_text("dict")
        for block in text_dict["blocks"]:
            if "lines" not in block:
                continue
            for line in block["lines"]:
                for span in line["spans"]:
                    text = span["text"]
                    if not text.strip():
                        continue
                    x0, y0, x1, y1 = span["bbox"]
                    left = Emu(int(x0 / 72 * EMU_PER_INCH))
                    top = Emu(int(y0 / 72 * EMU_PER_INCH))
                    width = Emu(max(int((x1 - x0) / 72 * EMU_PER_INCH), 1))
                    # a little extra height so descenders/ascenders aren't clipped
                    height = Emu(max(int((y1 - y0) / 72 * EMU_PER_INCH * 1.4), 1))

                    box = slide.shapes.add_textbox(left, top, width, height)
                    tf = box.text_frame
                    tf.word_wrap = True
                    tf.margin_left = 0
                    tf.margin_right = 0
                    tf.margin_top = 0
                    tf.margin_bottom = 0

                    run = tf.paragraphs[0].add_run()
                    run.text = text
                    font = run.font
                    font.size = Pt(max(span["size"], 1))
                    flags = span["flags"]
                    font.bold = bool(flags & 16)
                    font.italic = bool(flags & 2)

                    color_int = span["color"]
                    font.color.rgb = RGBColor(
                        (color_int >> 16) & 255, (color_int >> 8) & 255, color_int & 255,
                    )

                    font_name = span["font"]
                    if "+" in font_name:  # strip a subset prefix like "ABCDEF+Helvetica"
                        font_name = font_name.split("+", 1)[1]
                    for suffix in ("-Bold", "-Italic", "-BoldItalic", ",Bold", ",Italic"):
                        font_name = font_name.replace(suffix, "")
                    font.name = font_name or "Calibri"